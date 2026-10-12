"""Open workbooks robustly (build brief 8.1). The original blob is never modified."""

from __future__ import annotations

import csv
import re
import shutil
import tempfile
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl

PIVOT_PART = re.compile(r"^xl/(pivotCache|pivotTables)/")
PIVOT_REL = re.compile(r"pivotCache|pivotTable", re.IGNORECASE)


class UnsupportedWorkbook(Exception):
    """Cannot be parsed here; the caller should create a review item and keep the original."""


@dataclass
class RawSheet:
    name: str
    visible: bool
    rows: list[list[Any]]  # trimmed to the effective width
    max_row: int
    max_col_effective: int


@dataclass
class OpenedWorkbook:
    sheets: list[RawSheet]
    opened_via: str = "plain"  # plain | sanitized | csv
    warnings: list[str] = field(default_factory=list)


def sanitize_pivots(src: Path, dest: Path) -> None:
    """Write a copy of an xlsx without pivot caches/tables (the real master crashes openpyxl)."""
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dest, "w", zipfile.ZIP_DEFLATED) as zout:
        removed = {i.filename for i in zin.infolist() if PIVOT_PART.match(i.filename)}
        for info in zin.infolist():
            name = info.filename
            if name in removed:
                continue
            data = zin.read(name)
            if name == "xl/workbook.xml":
                data = re.sub(rb"<pivotCaches>.*?</pivotCaches>", b"", data, flags=re.DOTALL)
                data = re.sub(rb"<pivotCaches\s*/>", b"", data)
            elif name.endswith(".rels"):
                data = re.sub(
                    rb"<Relationship\b[^>]*>",
                    lambda m: (
                        b""
                        if PIVOT_REL.search(m.group(0).decode("utf-8", "ignore"))
                        else m.group(0)
                    ),
                    data,
                )
            elif name == "[Content_Types].xml":
                data = re.sub(
                    rb"<Override\b[^>]*PartName=\"/xl/(pivotCache|pivotTables)/[^>]*>", b"", data
                )
            zout.writestr(info, data)


def _effective_rows(ws: Any, scan_cols: int) -> tuple[list[list[Any]], int]:
    width = 0
    rows: list[list[Any]] = []
    max_col = min(ws.max_column or 1, scan_cols)
    for row in ws.iter_rows(max_col=max_col, values_only=True):
        r = list(row)
        while r and (r[-1] is None or (isinstance(r[-1], str) and not r[-1].strip())):
            r.pop()
        width = max(width, len(r))
        rows.append(r)
    while rows and not rows[-1]:
        rows.pop()
    return [r + [None] * (width - len(r)) for r in rows], width


def _read(wb: Any, scan_cols: int) -> list[RawSheet]:
    out: list[RawSheet] = []
    for ws in wb.worksheets:
        rows, width = _effective_rows(ws, scan_cols)
        out.append(
            RawSheet(
                name=ws.title,
                visible=ws.sheet_state == "visible",
                rows=rows,
                max_row=len(rows),
                max_col_effective=width,
            )
        )
    return out


def open_workbook(
    path: Path,
    derived_dir: Path | None = None,
    scan_cols: int = 512,
    suffix: str | None = None,
) -> OpenedWorkbook:
    """`suffix` overrides the file extension (blob-store files have none)."""
    suffix = (suffix or path.suffix).lower()
    if suffix == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as fh:
            rows: list[list[Any]] = [list(r) for r in csv.reader(fh)]
        width = max((len(r) for r in rows), default=0)
        rows = [r + [None] * (width - len(r)) for r in rows]
        return OpenedWorkbook([RawSheet(path.stem, True, rows, len(rows), width)], "csv")
    if suffix == ".xls":
        try:
            import xlrd
        except ImportError as exc:
            raise UnsupportedWorkbook(".xls needs xlrd, which is not installed") from exc
        book = xlrd.open_workbook(str(path))
        sheets = []
        for sh in book.sheets():
            rows = [list(sh.row_values(i)) for i in range(sh.nrows)]
            sheets.append(RawSheet(sh.name, True, rows, sh.nrows, sh.ncols))
        return OpenedWorkbook(sheets, "plain")
    if suffix not in {".xlsx", ".xlsm"}:
        raise UnsupportedWorkbook(f"unsupported workbook type {suffix!r}")
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        return OpenedWorkbook(_read(wb, scan_cols), "plain")
    except Exception as first:  # noqa: BLE001 - openpyxl raises TypeError/KeyError/etc. on bad pivots
        warn = (
            f"plain open failed ({type(first).__name__}: {str(first)[:120]}); "
            "opened a sanitized copy"
        )
        tmp_dir = derived_dir or Path(tempfile.mkdtemp(prefix="wcm-san-"))
        tmp_dir.mkdir(parents=True, exist_ok=True)
        copy = tmp_dir / f"sanitized-{path.stem}.xlsx"
        try:
            sanitize_pivots(path, copy)
            wb = openpyxl.load_workbook(copy, data_only=True)
        except Exception as second:  # noqa: BLE001
            raise UnsupportedWorkbook(f"cannot open even after sanitizing: {second}") from second
        res = OpenedWorkbook(_read(wb, scan_cols), "sanitized", [warn])
        if derived_dir is None:
            shutil.rmtree(tmp_dir, ignore_errors=True)
        return res
