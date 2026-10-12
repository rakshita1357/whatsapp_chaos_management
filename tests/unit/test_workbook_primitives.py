import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any

import pytest

from app.core.config import ROOT
from app.workbooks import open as wb_open
from app.workbooks.mapping import map_headers
from app.workbooks.normalize import division_key, normalize_cell
from app.workbooks.open import RawSheet, open_workbook, sanitize_pivots
from app.workbooks.roles import family_key, family_stem, infer_role
from app.workbooks.settings import load_workbook_config
from app.workbooks.sheets import classify_sheet
from tests.xlsx_helpers import HEADERS, add_fake_pivot, make_xlsx, row

CFG = load_workbook_config(ROOT / "config")


def test_open_plain_and_hidden_sheet_effective_width(tmp_path: Path) -> None:
    p = make_xlsx(
        tmp_path / "a.xlsx",
        {
            "Dealer List": [HEADERS, row(1, "3007530", "3007530-X", "Gaya", "Ready To Deploy")],
            "Hidden": [["a", "b"]],
        },
        hidden=("Hidden",),
    )
    import openpyxl

    wb = openpyxl.load_workbook(p)
    wb["Hidden"]["XFD1"] = "far away"  # INC-22: a hidden sheet reporting 16379+ columns
    wb.save(p)
    res = open_workbook(p, tmp_path / "d")
    assert res.opened_via == "plain"
    by = {s.name: s for s in res.sheets}
    assert by["Hidden"].visible is False
    assert by["Hidden"].max_col_effective <= 512
    assert by["Dealer List"].max_col_effective == len(HEADERS)


def test_sanitize_removes_pivot_parts_rels_and_caches(tmp_path: Path) -> None:
    base = make_xlsx(tmp_path / "a.xlsx", {"S": [HEADERS]})
    piv = add_fake_pivot(base, tmp_path / "p.xlsx")
    out = tmp_path / "clean.xlsx"
    sanitize_pivots(piv, out)
    with zipfile.ZipFile(out) as z:
        names = z.namelist()
        assert not any("pivot" in n.lower() for n in names)
        assert b"pivotCache" not in z.read("xl/workbook.xml")
        assert b"pivot" not in z.read("xl/_rels/workbook.xml.rels").lower()
        assert b"pivot" not in z.read("[Content_Types].xml").lower()
    assert open_workbook(out).sheets[0].name == "S"
    with zipfile.ZipFile(piv) as z:  # original untouched
        assert any("pivotCache" in n for n in z.namelist())


def test_open_falls_back_to_sanitized_copy(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    base = make_xlsx(tmp_path / "a.xlsx", {"S": [HEADERS]})
    piv = add_fake_pivot(base, tmp_path / "p.xlsx")
    real = wb_open.openpyxl.load_workbook
    calls: list[Path] = []

    def flaky(path: Any, *a: Any, **k: Any) -> Any:
        calls.append(Path(path))
        if Path(path) == piv:
            raise TypeError("pivot cache crash")  # what the real master does
        return real(path, *a, **k)

    monkeypatch.setattr(wb_open.openpyxl, "load_workbook", flaky)
    res = open_workbook(piv, tmp_path / "d")
    assert res.opened_via == "sanitized" and "TypeError" in res.warnings[0]
    assert len(calls) == 2 and calls[1] != piv


def test_csv_and_unsupported(tmp_path: Path) -> None:
    c = tmp_path / "x.csv"
    c.write_text("Dealer Code,City\n3007530,Gaya\n", encoding="utf-8")
    r = open_workbook(c)
    assert r.opened_via == "csv" and r.sheets[0].rows[1] == ["3007530", "Gaya"]
    with pytest.raises(wb_open.UnsupportedWorkbook):
        open_workbook(tmp_path / "x.docx")


def _sheet(rows: list[list[Any]], name: str = "S") -> RawSheet:
    w = max(len(r) for r in rows)
    return RawSheet(name, True, [r + [None] * (w - len(r)) for r in rows], len(rows), w)


def test_sheet_kinds() -> None:
    data = _sheet([["Report"], [None], HEADERS, [1, "3007530"]])
    info = classify_sheet(data, CFG)
    assert info.kind == "data_table" and info.header_row == 3
    pivot = _sheet([["Row Labels", "Count of Division"], ["x", 3], ["Grand Total", 3]])
    assert classify_sheet(pivot, CFG).kind == "pivot_summary"
    assert classify_sheet(_sheet([["notes", "misc"]]), CFG).kind == "helper_or_empty"


def test_header_mapping_exact_fuzzy_unknown() -> None:
    cols = map_headers(
        ["Dealer Code", "Pincode ", "Divison", "Mystery Column", "Dealer Code", None], CFG
    )
    by = {c.header_raw: c for c in cols}
    assert by["Dealer Code"].method == "exact"
    assert by["Pincode"].canonical == "pin" and by["Pincode"].method == "exact"
    assert by["Divison"].method == "fuzzy" and by["Divison"].canonical == "division_code"
    assert by["Mystery Column"].method == "unknown" and by["Mystery Column"].canonical is None
    assert by["Dealer Code#2"].canonical is None  # duplicate header never maps twice


def test_normalization_rules() -> None:
    n = lambda v, f, role="ANY": normalize_cell(v, f, role, CFG)  # noqa: E731
    assert n(datetime(2026, 10, 9, 5, 0), "site_ready_date").normalized == "2026-10-09"
    assert n("09/10/2026", "site_ready_date").normalized == "2026-10-09"  # D/M/Y
    assert n("1-3-26", "site_ready_date").normalized == "2026-03-01"
    bad = n("10-10-206", "site_ready_date")  # INC-44
    assert bad.normalized is None and bad.parse_status == "invalid"
    free = n("Pending", "site_ready_date")
    assert free.normalized is None and free.normalized_text == "Pending"
    assert n(" 300A490 ", "dealer_code").normalized == "300A490"
    assert n(3007530.0, "dealer_code").normalized == "3007530"
    assert n(9.88e9, "contact_phone").normalized == "9880000000"
    assert n("80 00-01", "pin").normalized == "800001"
    assert n("Ready To Deploy", "readiness").normalized == "ready"
    um = n("Maybe", "readiness")
    assert um.normalized is None and um.parse_status == "unmapped"
    assert n(7, "sr_no").normalized is None  # positional: never normalized into a key
    assert division_key(" 300A490 -Sv &Pa") == "300a490-sv&pa"


def test_roles_and_family_keys() -> None:
    assert infer_role(["Implementation Plan"], 10, [], False, CFG) == "IMPLEMENTATION_PLAN"
    assert infer_role(["x"], 3, ["ready", "ready", "ready"], False, CFG) == "READY_LIST"
    assert infer_role(["x"], 1200, ["ready"] * 600 + ["not_ready"] * 600, False, CFG) == "MASTER"
    assert infer_role(["x"], 50, ["ready", "not_ready"] * 25, False, CFG) == "UNKNOWN"
    assert infer_role(["p"], 0, [], True, CFG) == "SUMMARY"
    stems = {
        family_stem(n)
        for n in (
            "TMS_Delight_09-10-2026.xlsx",
            "TMS_Delight_10-10-2026-1.xlsx",
            "TMS Delight (1).xlsx",
            "TMS_Delight_20261009.xlsx",
        )
    }
    assert stems == {"tms delight"}
    canon = ["dealer_code", "city", None]
    assert family_key("A_01-01-2026.xlsx", "MASTER", canon) == family_key(
        "A-1.xlsx", "MASTER", canon
    )
    assert family_key("A.xlsx", "MASTER", canon) != family_key("A.xlsx", "READY_LIST", canon)
